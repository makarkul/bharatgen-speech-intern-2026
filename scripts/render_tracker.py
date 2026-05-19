"""Render the xlsx tracker into GitHub-friendly views.

Source of truth: docs/bharatgen_intern_tracker.xlsx
Generated:
  - docs/tracker.csv  (the 98-row Tracker sheet; GitHub auto-renders CSV as a sortable table)
  - docs/TRACKER.md   (Overview + Deliverables + Cadence + Success Criteria + Anti-Patterns
                       + Glossary + Learning Resources, plus summary stats)

Usage: python scripts/render_tracker.py

A GitHub Action at .github/workflows/render-tracker.yml runs this script
automatically when the xlsx (or this script) is pushed, and commits the
regenerated views back to the branch. Running locally is only needed if
you want to preview before pushing.
"""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "docs" / "bharatgen_intern_tracker.xlsx"
CSV_OUT = ROOT / "docs" / "tracker.csv"
MD_OUT = ROOT / "docs" / "TRACKER.md"

STATUS_EMOJI = {
    "Done": "✅",
    "In Progress": "🟡",
    "Blocked": "🔴",
    "Not Started": "⬜",
}


def cells(row):
    return ["" if v is None else str(v).strip() for v in row]


def md_escape(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ")


def md_table(header, rows):
    out = ["| " + " | ".join(md_escape(h) for h in header) + " |",
           "| " + " | ".join("---" for _ in header) + " |"]
    for r in rows:
        out.append("| " + " | ".join(md_escape(c) for c in r) + " |")
    return "\n".join(out)


def render_csv(wb):
    ws = wb["Tracker"]
    rows = [cells(r) for r in ws.iter_rows(values_only=True)]
    with CSV_OUT.open("w", newline="") as f:
        csv.writer(f).writerows(rows)


def render_md(wb):
    overview = list(wb["Overview"].iter_rows(values_only=True))

    def overview_get(label):
        for row in overview:
            if row and row[0] == label:
                return row[1]
        return None

    tracker = wb["Tracker"]
    header = [c.value for c in tracker[1]]
    status_idx = header.index("Status")
    statuses = Counter(
        (r[status_idx] or "Not Started")
        for r in tracker.iter_rows(min_row=2, values_only=True)
    )
    total = sum(statuses.values())

    parts = []
    parts.append("# BharatGen Speech Intern — Tracker\n")
    parts.append(
        "Source of truth: [`bharatgen_intern_tracker.xlsx`](bharatgen_intern_tracker.xlsx). "
        "This file is **generated** by `scripts/render_tracker.py`; edit the xlsx and rerun.\n"
    )

    parts.append("## Status summary\n")
    summary_rows = [
        ["Total items", str(total)],
        *[[f"{STATUS_EMOJI.get(k, '')} {k}", str(v)] for k, v in statuses.most_common()],
        ["% Complete", f"{(statuses.get('Done', 0) / total * 100):.1f}%" if total else "—"],
    ]
    parts.append(md_table(["Metric", "Count"], summary_rows) + "\n")

    parts.append("## Project setup\n")
    setup_rows = []
    for label in ("Intern name", "Mentor name", "Start date (Week 0 Day 1)",
                  "End date (computed)", "Repo URL", "GPU environment",
                  "HF token configured?"):
        val = overview_get(label)
        setup_rows.append([label, "" if val is None else str(val)])
    parts.append(md_table(["Field", "Value"], setup_rows) + "\n")

    parts.append(
        "## Full tracker (98 items)\n\n"
        "GitHub renders CSV files as a sortable, searchable table — open "
        "[`tracker.csv`](tracker.csv).\n"
    )

    # Per-week status table
    parts.append("## Per-week status\n")
    by_week: dict[str, Counter] = {}
    for r in tracker.iter_rows(min_row=2, values_only=True):
        week = r[0] or "—"
        status = r[status_idx] or "Not Started"
        by_week.setdefault(week, Counter())[status] += 1
    week_rows = []
    for week, counts in by_week.items():
        wk_total = sum(counts.values())
        done = counts.get("Done", 0)
        bar = f"{done}/{wk_total}"
        pct = f"{(done / wk_total * 100):.0f}%" if wk_total else "—"
        week_rows.append([week, bar, pct,
                          str(counts.get("In Progress", 0)),
                          str(counts.get("Blocked", 0)),
                          str(counts.get("Not Started", 0))])
    parts.append(md_table(
        ["Week", "Done", "%", "🟡 In Progress", "🔴 Blocked", "⬜ Not Started"],
        week_rows,
    ) + "\n")

    # Remaining sheets
    sheet_titles = {
        "Deliverables": "Weekly deliverables",
        "Cadence": "Weekly cadence",
        "Success Criteria": "Week-7 success criteria",
        "Anti-Patterns": "Anti-patterns",
        "Glossary": "Glossary",
        "Learning Resources": "Learning resources",
    }
    for sheet_name, title in sheet_titles.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = [cells(r) for r in ws.iter_rows(values_only=True)]
        # Strip leading title/blurb rows; find the header row (first row with >=2 non-empty cells
        # *and* second row also non-empty — i.e., looks like a header).
        header_idx = 0
        for i, r in enumerate(rows):
            non_empty = sum(1 for c in r if c)
            if non_empty >= 2 and i + 1 < len(rows) and any(rows[i + 1]):
                header_idx = i
                break
        table_rows = [r for r in rows[header_idx + 1:] if any(c for c in r)]
        if not table_rows:
            continue
        parts.append(f"## {title}\n")
        parts.append(md_table(rows[header_idx], table_rows) + "\n")

    MD_OUT.write_text("\n".join(parts))


def main():
    wb = load_workbook(XLSX, data_only=True)
    render_csv(wb)
    render_md(wb)
    print(f"Wrote {CSV_OUT.relative_to(ROOT)} and {MD_OUT.relative_to(ROOT)}.")


if __name__ == "__main__":
    main()
